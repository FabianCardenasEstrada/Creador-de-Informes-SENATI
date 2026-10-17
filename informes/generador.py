"""Llena las plantillas oficiales (empresa y seminario) con los datos del diario."""
import copy
import shutil
import subprocess
from datetime import timedelta
from pathlib import Path

import docx
from docx.shared import Cm
from lxml import etree

from . import docx_utils as X
from .docx_utils import W
from .diagrama import dibujar
from .util import (RAIZ, ruta, f_corta, f_larga, nombre_dia, fmt_horas,
                   semanas_lun_sab, trozos_rotacion)

SZ = 18          # tamaño (half-points) del texto dentro de cuadros: 9 pt
SZ_TXT = 20      # tamaño del texto de narrativa: 10 pt


class Generador:
    def __init__(self, cfg, pea):
        self.cfg = cfg
        self.pea = pea

    # ================================================================ helpers de búsqueda
    @staticmethod
    def _tabla(doc, inicio_texto, nth=0):
        hallados = [t for t in doc.tables if t.rows[0].cells[0].text.strip().startswith(inicio_texto)]
        return hallados[nth] if len(hallados) > nth else None

    @staticmethod
    def _tabla_contiene(doc, texto):
        return [t for t in doc.tables if texto in t.rows[0].cells[0].text]

    @staticmethod
    def _semanales(doc):
        return [t for t in doc.tables if t.rows[0].cells[0].text.strip() == "DÍA"
                and len(t.rows) == 8 and len(t.columns) == 3]

    # ================================================================ utilidades de formato
    @staticmethod
    def _titulo(doc, inicio):
        """Busca un título de la plantilla y le quita las tabulaciones que dibujan líneas punteadas."""
        p = X.buscar_parrafo(doc, inicio)
        for r in p.findall(W("r")):
            if r.find(W("tab")) is not None and r.find(W("t")) is None:
                p.remove(r)
        return p

    @staticmethod
    def _mantener_junta(tabla):
        filas = tabla._tbl.findall(W("tr"))
        for i, tr in enumerate(filas):
            trpr = tr.find(W("trPr"))
            if trpr is None:
                trpr = etree.Element(W("trPr"))
                tr.insert(0, trpr)
            if trpr.find(W("cantSplit")) is None:
                trpr.insert(0, etree.Element(W("cantSplit")))
            if i < len(filas) - 1:
                for p in tr.iter(W("p")):
                    X.keep_next(p)

    # ================================================================ portada
    def _portada(self, doc, tipo):
        e, s = self.cfg["estudiante"], self.cfg["semestre"]

        def etiqueta(inicio, valor):
            p = X.buscar_parrafo(doc, inicio)
            runs = p.findall(W("r"))
            primero = next(r for r in runs if r.find(W("t")) is not None)
            primero.addnext(X.mk_run(valor, sz=24))

        etiqueta("CFP/UCP/ESCUELA", e["cfp"])
        etiqueta("ESTUDIANTE", e["nombre"])
        etiqueta("CARRERA", e["carrera"])
        etiqueta("INSTRUCTOR", e["instructor"])

        # ID / BLOQUE y SEMESTRE / DEL / AL: se reconstruyen con las tabulaciones de la plantilla
        def reconstruir(inicio, partes):
            p = X.buscar_parrafo(doc, inicio)
            for r in p.findall(W("r")):
                p.remove(r)
            for texto, es_tab in partes:
                p.append(X.mk_run(texto, sz=24, tab=es_tab))

        reconstruir("ID:", [("ID: ", False), (e["id"], False), ("", True),
                            (" BLOQUE: ", False), (e["bloque"], False), ("", True)])
        reconstruir("SEMESTRE", [("SEMESTRE: ", False), (s["numero"], False), ("", True),
                                 (" DEL: ", False), (s["del"], False), ("", True),
                                 (" AL: ", False), (s["al"], False), ("", True)])

        # Dirección zonal (cuadro de texto flotante)
        p = X.buscar_parrafo(doc, "DIRECCIÓN ZONAL")
        for t in p.iter(W("t")):
            if t.text and set(t.text) == {"_"}:
                t.text = e["direccion_zonal"]

        # Código del informe
        codigo = self.cfg["empresa" if tipo == "empresa" else "seminario"].get("codigo_informe", "")
        if codigo and tipo == "empresa":
            p = X.buscar_parrafo(doc, "CÓDIGO N°")
            ts = [t for t in p.iter(W("t"))]
            ts[-1].text = f"CÓDIGO N° {codigo}"
            ts[-1].set("{http://www.w3.org/XML/1998/namespace}space", "preserve")

    # ================================================================ rotaciones
    def _rotaciones(self, doc, filas):
        t = self._tabla_contiene(doc, "ROTACIONES")[0]
        for i, fila in enumerate(filas):
            r = t.rows[3 + i]
            celdas = X.celdas_unicas(r)
            for c, txt in zip(celdas, fila):
                X.poner_en_celda(c._tc, [X.par_simple(txt, sz=SZ, jc="center", after=0)])

    # ================================================================ PEA
    def _pea(self, doc, estado, tipo):
        t = self._tabla(doc, "Nº")
        filas_datos = t.rows[2:]
        ops = sorted(estado)
        for i, fila in enumerate(filas_datos):
            if i >= len(ops):
                X.quitar(fila._tr)
                continue
            n = ops[i]
            cel = fila.cells
            nuevas = []
            for c in X.celdas_unicas(fila):
                nuevas.append(c)
            # columnas: 0 N°, 1 texto, 2..5 = marcas 1..4, 6 (y 7 en seminario)
            X.poner_en_celda(cel[0]._tc, [X.par_simple(f"{n:02d}", sz=SZ, jc="center", after=0)])
            X.poner_en_celda(cel[1]._tc, [X.par_simple(self.pea.texto(n), sz=SZ, jc="left", after=0)])
            for col in (2, 3, 4, 5):
                marca = "X" if (estado[n] == 2 and col == 3) or (estado[n] == 4 and col == 5) else ""
                X.poner_en_celda(cel[col]._tc, [X.par_simple(marca, sz=SZ, jc="center", after=0)])
            for col in range(6, len(cel)):
                X.poner_en_celda(cel[col]._tc, [X.par_simple("-", sz=SZ, jc="center", after=0)])

    # ================================================================ semanas
    def _contenido_dia(self, reg, tipo_doc):
        """Devuelve (lista de w:p, horas)."""
        if not reg:
            return [], 0
        tp = reg["tipo"]
        if tipo_doc == "seminario" and tp != "seminario":
            return [], 0
        etiquetas = {
            "empresa": self.cfg["empresa"]["etiqueta_dia"],
            "tecnologia": self.cfg["empresa"]["etiqueta_tecnologia"],
            "seminario": self.cfg["seminario"]["etiqueta_dia"],
            "falta": "SENATI:" if reg.get("lugar") == "senati" else "Empresa:",
            "feriado": "",
        }
        elems = []
        if etiquetas[tp]:
            elems.append(X.par_simple(etiquetas[tp], bold=True, sz=16, jc="center", after=40))
        ops = reg.get("operaciones", [])
        if tp == "seminario":
            ops = sorted({o for ej in reg.get("ejercicios", []) for o in ej.get("operaciones", [])} | set(ops))
        if ops:
            elems.append(X.par_simple("Operaciones", bold=True, sz=SZ, jc="left", after=20))
            for n in ops:
                elems.append(X.par_vineta(self.pea.texto(n), sz=SZ))
            elems.append(X.par_simple("", sz=8, after=0))
        if tp == "seminario":
            for ej in reg.get("ejercicios", []):
                runs = [X.mk_run(ej["titulo"].rstrip(".") + ". ", bold=True, sz=SZ),
                        X.mk_run(ej.get("descripcion", ""), sz=SZ)]
                elems.append(X.mk_par(runs, jc="both", after=60))
        elif tp == "falta":
            just = "justificada" if reg.get("asistencia") != "falta_injustificada" else "injustificada"
            motivo = reg.get("motivo", "").strip()
            donde = "a SENATI" if reg.get("lugar") == "senati" else "a la empresa"
            txt = f"Se registró inasistencia {just} {donde}" + (f" por motivo de {motivo}." if motivo else ".")
            elems.append(X.par_simple(txt, sz=SZ, jc="center"))
        elif tp == "feriado":
            elems.append(X.par_simple(reg.get("descripcion") or "Día feriado / sin actividad.", sz=SZ, jc="center"))
        else:
            if reg.get("descripcion"):
                elems.append(X.par_simple(reg["descripcion"], sz=SZ, jc="both"))
        return elems, float(reg.get("horas", 0) or 0)

    def _llenar_semana(self, tabla, lunes, regs, tipo_doc, inicio, fin):
        total = 0.0
        for i in range(6):
            fecha = lunes + timedelta(days=i)
            celdas = X.celdas_unicas(tabla.rows[i + 1])
            X.poner_en_celda(celdas[0]._tc, [
                X.par_simple(nombre_dia(fecha), bold=True, sz=SZ, jc="center", after=0),
                X.par_simple(f_larga(fecha), sz=16, jc="center", after=0)])
            reg = regs.get(fecha) if inicio <= fecha <= fin else None
            elems, horas = self._contenido_dia(reg, tipo_doc)
            X.poner_en_celda(celdas[1]._tc, elems)
            X.poner_en_celda(celdas[2]._tc,
                             [X.par_simple(fmt_horas(horas), sz=SZ, jc="center", after=0)] if elems else [])
            total += horas
        self._mantener_junta(tabla)
        ult = X.celdas_unicas(tabla.rows[7])[-1]
        X.poner_en_celda(ult._tc, [X.par_simple(fmt_horas(total), bold=True, sz=SZ, jc="center", after=0)])
        return total

    # ================================================================ narrativa
    @staticmethod
    def _antes_de_tabla(tabla_el, elems):
        for e in elems:
            tabla_el.addprevious(e)

    def _narrativa_empresa(self, doc, d):
        narr, cierre = d["narr"], d["cierre"]
        p = self._titulo(doc, "Tarea más significativa durante")
        X.rellenar_despues(p, [
            X.mk_par([X.mk_run("Tarea: ", bold=True, sz=SZ_TXT), X.mk_run(narr["tarea"].rstrip(".") + ".", sz=SZ_TXT)],
                     jc="both", after=60),
            X.par_simple(narr["por_que"], sz=SZ_TXT, jc="both", after=120)])
        p = self._titulo(doc, "Descripción del proceso")
        X.rellenar_despues(p, [X.par_vineta(s, sz=SZ_TXT) for s in narr["proceso"]] + [X.par_simple("", sz=8)])
        p = self._titulo(doc, "Máquinas, equipos")
        X.rellenar_despues(p, [X.par_vineta(h, sz=SZ_TXT, caracter="-") for h in narr["herramientas"]] + [X.par_simple("", sz=8)])
        p = self._titulo(doc, "Seguridad e higiene")
        X.rellenar_despues(p, [X.par_simple(d["seguridad"], sz=SZ_TXT, jc="both", after=120)])
        p = self._titulo(doc, "Resultados de la ejecución")
        elems = [X.par_simple(cierre["resultados"], sz=SZ_TXT, jc="both", after=80)]
        if cierre["recomendaciones"]:
            elems.append(X.par_simple("Recomendaciones:", bold=True, sz=SZ_TXT, jc="left", after=20))
            elems += [X.par_vineta(r, sz=SZ_TXT) for r in cierre["recomendaciones"]]
        X.rellenar_despues(p, elems)
        t = self._tabla_contiene(doc, "HACER ESQUEMA")[0]
        t._tbl.addprevious(X.par_salto_pagina())

    def _narrativa_seminario(self, doc, d):
        narr = d["narr"]
        p = self._titulo(doc, "Tarea más significativa")
        r = p.findall(W("r"))[-1] if p.findall(W("r")) else None
        run = X.mk_run(" " + narr["tarea"].rstrip(".") + ".", sz=SZ_TXT)
        (r.addnext(run) if r is not None else p.append(run))
        p = self._titulo(doc, "Descripción del proceso")
        elems = [X.par_simple(narr["por_que"], sz=SZ_TXT, jc="both", after=100)]
        if narr["proceso"]:
            elems.append(X.par_simple("Secuencia de ejecución:", bold=True, sz=SZ_TXT, jc="left", after=20))
            elems += [X.par_vineta(s, sz=SZ_TXT) for s in narr["proceso"]]
        if narr["herramientas"]:
            elems.append(X.par_simple("", sz=8))
            elems.append(X.mk_par([X.mk_run("Herramientas utilizadas: ", bold=True, sz=SZ_TXT),
                                   X.mk_run(", ".join(narr["herramientas"]) + ".", sz=SZ_TXT)], jc="both", after=60))
        X.rellenar_despues(p, elems)
        t = self._tabla_contiene(doc, "HACER ESQUEMA")[0]
        t._tbl.addprevious(X.par_salto_pagina())

    def _diagrama(self, doc, nodos, titulo):
        t = self._tabla_contiene(doc, "HACER ESQUEMA")[0]
        fila = t.rows[1]
        tc = fila.cells[0]._tc
        ancho = X.ancho_celda_twips(tc) / 567.0 - 0.5
        alto_tw = X.alto_fila_twips(fila._tr)
        alto = (alto_tw / 567.0 - 0.4) if alto_tw else 9.5
        alto = max(6.0, min(alto, 22.0))
        png = RAIZ / "salida" / "_diagrama.png"
        dibujar(nodos, str(png), ancho_cm=ancho, alto_cm=alto)
        celda = fila.cells[0]
        X.limpiar_celda(tc)
        tc.append(X.mk_par([], jc="center"))
        celda.paragraphs[0].alignment = 1
        celda.paragraphs[0].add_run().add_picture(str(png), width=Cm(ancho), height=Cm(alto))

    # ================================================================ recorte de bloques semanales
    @staticmethod
    def _colapsar_saltos(doc):
        """Evita páginas en blanco por saltos de página consecutivos."""
        def es_salto(p):
            return p.tag == W("p") and not X.texto_de(p).strip() and any(
                b.get(W("type")) == "page" for b in p.iter(W("br")))
        cuerpo = list(doc.element.body.iterchildren())
        for i, el in enumerate(cuerpo):
            if not es_salto(el):
                continue
            j = i + 1
            while j < len(cuerpo) and X.es_vacio(cuerpo[j]):
                j += 1
            if j < len(cuerpo) and es_salto(cuerpo[j]):
                X.quitar(el)

    def _quitar_tabla_semanal(self, tabla, quitar_encabezado=False):
        el = tabla._tbl
        if quitar_encabezado:
            prev = el.getprevious()
            while prev is not None:
                anterior = prev.getprevious()
                txt = X.texto_de(prev).strip() if prev.tag == W("p") else ""
                X.quitar(prev)
                if txt.startswith("INFORME SEMANAL"):
                    break
                prev = anterior
        for v in X.vacios_siguientes(el):
            X.quitar(v)
        X.quitar(el)

    # ================================================================ empresa
    def empresa(self, d, salida):
        doc = docx.Document(str(ruta("plantillas/empresa.docx")))
        inicio, fin = d["inicio"], d["fin"]
        self._portada(doc, "empresa")
        area = self.cfg["empresa"]["area_rotacion"]
        self._rotaciones(doc, [(area, f_corta(a), f_corta(b), "1") for a, b in trozos_rotacion(inicio, fin)])
        self._pea(doc, d["estado"], "empresa")

        regs = {f: r for f, r in d["registros"]}
        lunes_l = semanas_lun_sab(inicio, fin)
        tablas = self._semanales(doc)
        while len(tablas) < len(lunes_l):          # clonar si hacen falta más semanas
            nueva = copy.deepcopy(tablas[-1]._tbl)
            tablas[-1]._tbl.addnext(nueva)
            tablas[-1]._tbl.addnext(X.par_simple("", sz=8))
            tablas = self._semanales(doc)
        for t, lunes in zip(tablas, lunes_l):
            self._llenar_semana(t, lunes, regs, "empresa", inicio, fin)
        for t in tablas[len(lunes_l):]:
            self._quitar_tabla_semanal(t)

        self._narrativa_empresa(doc, d)
        self._diagrama(doc, d["narr"]["diagrama"], None)
        self._evaluacion(doc, d)
        self._colapsar_saltos(doc)
        self._limpiar_final(doc)
        doc.save(str(salida))
        return salida

    def _evaluacion(self, doc, d):
        t = self._tabla_contiene(doc, "EVALUACIÓN DEL INFORME")[0]
        # Nota del instructor
        if d.get("nota"):
            X.texto_en_formas(t.rows[0].cells[0]._tc, [d["nota"]], sz=18)
        # Observaciones del monitor
        if d.get("comentario_monitora"):
            tc = t.rows[1].cells[0]._tc
            ps = tc.findall(W("p"))
            titulo = ps[0]
            for p in ps[1:]:
                tc.remove(p)
            tc.append(X.par_simple(d["comentario_monitora"], sz=SZ, jc="both", before=60))
        # Firmas
        fila = t.rows[3]
        firmas = self.cfg["firmas"]
        for celda, clave, activo in ((fila.cells[0], "alumno", True),
                                     (fila.cells[1], "monitora", d.get("firmar_monitora", False))):
            if not activo:
                continue
            ruta_f = ruta(firmas[clave])
            if not ruta_f.exists():
                continue
            X.limpiar_celda(celda._tc)
            celda._tc.append(X.mk_par([], jc="center"))
            celda.paragraphs[0].alignment = 1
            celda.paragraphs[0].add_run().add_picture(str(ruta_f), height=Cm(1.5))

    def _limpiar_final(self, doc):
        """Quita párrafos vacíos sobrantes del final para no generar una hoja en blanco."""
        cuerpo = doc.element.body
        hijos = [h for h in cuerpo.iterchildren() if h.tag != W("sectPr")]
        while hijos and X.es_vacio(hijos[-1]):
            X.quitar(hijos.pop())

    # ================================================================ seminario
    def seminario(self, d, salida):
        doc = docx.Document(str(ruta("plantillas/seminario.docx")))
        inicio, fin = d["inicio"], d["fin"]
        self._portada(doc, "seminario")
        lunes_l = semanas_lun_sab(inicio, fin)
        if len(lunes_l) > 4:
            raise ValueError("La plantilla del seminario admite hasta 4 semanas por informe.")
        self._rotaciones(doc, [(self.cfg["seminario"]["area_rotacion"], f_corta(inicio), f_corta(fin), str(len(lunes_l)))])
        self._pea(doc, d["estado"], "seminario")

        regs = {f: r for f, r in d["registros"]}
        tablas = self._semanales(doc)
        for k, t in enumerate(tablas):
            if k < len(lunes_l):
                self._llenar_semana(t, lunes_l[k], regs, "seminario", inicio, fin)
                self._encabezado_semana(t, k + 1, lunes_l[k])
            else:
                self._quitar_tabla_semanal(t, quitar_encabezado=True)

        self._narrativa_seminario(doc, d)
        self._diagrama(doc, d["narr"]["diagrama"], None)
        self._asistencia(doc, lunes_l, regs, inicio, fin)
        self._colapsar_saltos(doc)
        self._limpiar_final(doc)
        doc.save(str(salida))
        return salida

    def _encabezado_semana(self, tabla, numero, lunes):
        s = self.cfg["semestre"]
        prev = tabla._tbl.getprevious()
        while prev is not None and not (prev.tag == W("p") and "SEMANA" in X.texto_de(prev)):
            prev = prev.getprevious()
        if prev is None:
            return
        for r in prev.findall(W("r")):
            prev.remove(r)
        # mantener título + encabezado + tabla en la misma página
        q = tabla._tbl.getprevious()
        while q is not None:
            if q.tag == W("p"):
                X.keep_next(q)
                if X.texto_de(q).strip().startswith("INFORME SEMANAL"):
                    break
            q = q.getprevious()
        sab = lunes + timedelta(days=5)
        for texto, tab in ((f"{s['numero']} SEMESTRE", False), ("", True),
                           (f"SEMANA N° {numero}", False), ("     ", False),
                           (f"DEL {f_corta(lunes)} AL {f_corta(sab)} DEL {s['anio']}", False)):
            prev.append(X.mk_run(texto, sz=22, tab=tab))

    def _asistencia(self, doc, lunes_l, regs, inicio, fin):
        t = self._tabla_contiene(doc, "AUTOCONTROL")[0]
        colores = self.cfg["colores_asistencia"]
        filas = list(t.rows)
        for k in range(4):
            base = 4 * k
            if k >= len(lunes_l):
                for r in filas[base:base + 4]:
                    X.quitar(r._tr)
                continue
            fila_mt = X.celdas_unicas(filas[base + 2])
            cuentas = {"asistio": 0, "inj": 0, "just": 0}
            for i in range(6):
                fecha = lunes_l[k] + timedelta(days=i)
                reg = regs.get(fecha) if inicio <= fecha <= fin else None
                if not reg or not reg.get("asistencia") or reg["tipo"] not in ("tecnologia", "seminario", "falta"):
                    continue
                if reg["tipo"] == "falta" and reg.get("lugar") != "senati":
                    continue
                est = reg["asistencia"]
                turnos = [reg["turno"]] if reg.get("turno") in ("M", "T") else ["M", "T"]
                if est in ("temprano", "tarde"):
                    cuentas["asistio"] += 1
                elif est == "falta_injustificada":
                    cuentas["inj"] += 1
                elif est == "falta_justificada":
                    cuentas["just"] += 1
                for tu in turnos:
                    X.pintar(fila_mt[i * 2 + (0 if tu == "M" else 1)]._tc, colores[est])
            leyenda = X.celdas_unicas(filas[base + 3])[0]._tc

            def clave(a, _ns=X):
                try:
                    v = int(a.find(".//{%s}positionV/{%s}posOffset" % (X.NS_WP, X.NS_WP)).text)
                    h = int(a.find(".//{%s}positionH/{%s}posOffset" % (X.NS_WP, X.NS_WP)).text)
                except Exception:
                    v = h = 0
                return (0, 0) if h < 3000000 else (1, v)
            X.texto_en_formas(leyenda, [cuentas["asistio"] or "", cuentas["inj"] or "", cuentas["just"] or ""],
                              sz=16, clave=clave)


# ====================================================================== PDF
def convertir_pdf(docx_path):
    """Intenta crear el PDF con Word (docx2pdf) o con LibreOffice. Devuelve la ruta o None."""
    docx_path = Path(docx_path)
    try:
        from docx2pdf import convert  # requiere Microsoft Word en Windows/Mac
        convert(str(docx_path), str(docx_path.with_suffix(".pdf")))
        if docx_path.with_suffix(".pdf").exists():
            return docx_path.with_suffix(".pdf")
    except Exception:
        pass
    for exe in ("soffice", "libreoffice"):
        if shutil.which(exe):
            try:
                subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir",
                                str(docx_path.parent), str(docx_path)],
                               check=True, capture_output=True, timeout=180)
                if docx_path.with_suffix(".pdf").exists():
                    return docx_path.with_suffix(".pdf")
            except Exception:
                pass
    return None
